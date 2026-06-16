#!/usr/bin/env python3
"""
사내 업무채널 사용 방법 인수인계 자료
슬랙 / 줌 / 클로바노트 / 드롭박스
"""

from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.lib.units import cm
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle,
    HRFlowable, PageBreak, KeepTogether
)
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from PIL import Image, ImageDraw, ImageFont
import os

FONT_DIR = "/usr/share/fonts/truetype/nanum"
REGULAR  = os.path.join(FONT_DIR, "NanumGothic.ttf")
BOLD     = os.path.join(FONT_DIR, "NanumGothicBold.ttf")
SQUARE_B = os.path.join(FONT_DIR, "NanumSquareB.ttf")

pdfmetrics.registerFont(TTFont("NanumGothic",     REGULAR))
pdfmetrics.registerFont(TTFont("NanumGothicBold", BOLD))
pdfmetrics.registerFont(TTFont("NanumSquareB",    SQUARE_B))

OUT_DIR = "/home/user/project-1/handover_docs"
os.makedirs(OUT_DIR, exist_ok=True)

# 채널별 브랜드 컬러
C_SLACK   = colors.HexColor("#4a154b")   # 슬랙 보라
C_ZOOM    = colors.HexColor("#0b5cff")   # 줌 파랑
C_CLOVA   = colors.HexColor("#03c75a")   # 클로바노트 녹색
C_DROP    = colors.HexColor("#0061ff")   # 드롭박스 파랑
C_HEADER  = colors.HexColor("#2d2d2d")   # 표지/구분 헤더

# ══════════════════════════════════════════════════════════════
# 예시 이미지 생성
# ══════════════════════════════════════════════════════════════

def make_img(title, lines, filename, width=760, height=380,
             header_bg=(74, 21, 75)):
    img = Image.new("RGB", (width, height), (248, 248, 252))
    draw = ImageDraw.Draw(img)
    try:
        fnt_t = ImageFont.truetype(BOLD,    20)
        fnt_b = ImageFont.truetype(REGULAR, 16)
        fnt_s = ImageFont.truetype(REGULAR, 13)
    except Exception:
        fnt_t = fnt_b = fnt_s = ImageFont.load_default()

    draw.rectangle([0, 0, width, 52], fill=header_bg)
    draw.text((18, 14), title, font=fnt_t, fill=(255, 255, 255))

    y = 70
    for line in lines:
        if line.startswith("##"):
            draw.rectangle([10, y-2, width-10, y+26], fill=(230, 225, 245))
            draw.text((18, y+2), line[2:].strip(), font=fnt_b, fill=header_bg)
            y += 36
        elif line.startswith(">>"):
            draw.polygon([(24, y+6),(32, y+11),(24, y+16)], fill=header_bg)
            draw.text((40, y+2), line[2:].strip(), font=fnt_b, fill=(40, 40, 40))
            y += 28
        elif line.startswith("--"):
            draw.line([28, y+10, width-28, y+10], fill=(200, 200, 220), width=1)
            y += 20
        elif line.startswith("**"):
            draw.text((24, y), line[2:].strip(), font=fnt_b, fill=(60, 60, 60))
            y += 26
        else:
            draw.text((30, y), line, font=fnt_s, fill=(90, 90, 90))
            y += 22

    draw.rectangle([0, 0, width-1, height-1], outline=(180, 180, 210), width=2)
    path = os.path.join(OUT_DIR, filename)
    img.save(path)
    return path


img_slack_ch = make_img(
    "슬랙 (Slack) – ONNAMU 사내 채널 목록",
    [
        "## 공개 채널 (#)",
        "  #0-소셜-전체소통방       전체 구성원 자유 소통",
        "  #0-주간회의                    주간 회의 자료 및 공유",
        "  #1-근무및휴가                근무·휴가 신청 및 공유",
        "  #1-대표님-지시사항        대표님 지시사항 전달",
        "  #공모사업                        공모사업 관련 정보 공유",
        "  #김보고서-알림              보고서 알림 봇 채널",
        "  #indonesia인니팀-행정   인도네시아팀 행정 채널",
        "--",
        "## 비공개 채널 (🔒) – 초대된 멤버만 접근",
        "  🔒 3language-translation     다국어 번역 채널",
        "  🔒 ai-에이전트-테스트      AI 에이전트 테스트",
        "  🔒 legal-review                    법무 검토",
        "  🔒 sns-content-design         SNS 콘텐츠 디자인",
        "  🔒 summary-visualization    요약·시각화",
        "  🔒 youtube-summary             유튜브 요약",
    ],
    "img_s01_slack_ch.png",
    height=440,
    header_bg=(74, 21, 75)
)

img_slack_use = make_img(
    "슬랙 (Slack) – 주요 기능",
    [
        "## 메시지 작성",
        ">> @ 멘션 : @이름 또는 @here / @channel 으로 알림 발송",
        ">> 스레드 : 메시지에 [답글 달기] → 대화 흐름 정리",
        ">> 이모지 반응 : 메시지에 마우스 오버 → 😊 클릭",
        "--",
        "## 파일 공유 및 검색",
        ">> 파일 첨부 : 메시지 입력창 왼쪽 📎 클릭 또는 드래그 앤 드롭",
        ">> 검색 : 상단 검색창에 키워드 입력 (채널·기간 필터 가능)",
        "--",
        "  ★ 중요 메시지 : 메시지 오버 → 북마크(🔖) 저장",
        "  ★ 알림 설정 : 프로필 > 환경설정 > 알림에서 조정",
    ],
    "img_s02_slack_use.png",
    header_bg=(74, 21, 75)
)

img_zoom = make_img(
    "줌 (Zoom) – 회의 진행",
    [
        "## 회의 시작 / 참여",
        ">> 새 회의 : Zoom 앱 실행 → [새 회의] 클릭",
        ">> 회의 예약 : [예약] → 날짜·시간 설정 → 초대링크 복사 후 공유",
        ">> 회의 참여 : 수신된 링크 클릭 또는 회의 ID 입력",
        "--",
        "## 회의 중 주요 기능",
        ">> 마이크 : 하단 🎤 클릭으로 음소거/해제",
        ">> 화면 공유 : 하단 [공유] → 공유할 화면/창 선택",
        ">> 채팅 : 하단 [채팅] → 참여자 전체 또는 개별 메시지",
        ">> 녹화 : [녹화] 클릭 → 로컬 또는 클라우드 저장",
        "--",
        "  ★ 회의 링크 : 슬랙 채널에 공유하여 참여자 안내",
    ],
    "img_s03_zoom.png",
    header_bg=(11, 92, 255)
)

img_clova = make_img(
    "클로바노트 (ClovaNote) – 회의 녹음 및 요약",
    [
        "## 회의 녹음 시작",
        ">> 클로바노트 앱(모바일) 또는 웹(clovanote.naver.com) 실행",
        ">> [새 노트] → [녹음 시작] → 회의 중 자동 음성 인식",
        ">> 녹음 종료 후 자동으로 텍스트 변환 및 화자 분리",
        "--",
        "## 회의록 활용",
        ">> 텍스트 변환본 검토 후 오류 수정",
        ">> AI 요약 기능으로 핵심 내용 자동 정리 확인",
        ">> [공유] → 링크 또는 텍스트 복사 → 슬랙 채널에 공유",
        "--",
        "  ★ 줌 회의 시 : 줌 녹화(mp3) 파일을 클로바노트에 업로드 후 변환 가능",
        "  ★ 회의록 보관 : 드롭박스 '회의록/YYYY-MM' 폴더에 저장",
    ],
    "img_s04_clova.png",
    header_bg=(0, 160, 70)
)

img_dropbox = make_img(
    "드롭박스 (Dropbox) – 파일 보관 및 공유",
    [
        "## 폴더 구조 (예시)",
        "  📁 공유폴더/",
        "     📁 계약서/          계약 관련 문서",
        "     📁 회의록/          월별 회의록",
        "     📁 세금계산서/   월별 세금계산서 PDF",
        "     📁 인수인계/       업무 인수인계 자료",
        "--",
        "## 파일 공유 방법",
        ">> 파일/폴더 우클릭 → [링크 복사] → 슬랙 등에 링크 붙여넣기",
        ">> 팀 폴더 : 초대된 멤버 자동 동기화",
        "--",
        "  ★ 중요 문서는 반드시 드롭박스에 보관 (로컬 저장 금지)",
        "  ★ 파일명 규칙 : YYYY-MM-DD_문서명.확장자",
    ],
    "img_s05_dropbox.png",
    header_bg=(0, 97, 255)
)


# ══════════════════════════════════════════════════════════════
# PDF 생성
# ══════════════════════════════════════════════════════════════

def build_pdf(out_path):
    doc = SimpleDocTemplate(
        out_path, pagesize=A4,
        topMargin=2*cm, bottomMargin=2*cm,
        leftMargin=2.2*cm, rightMargin=2.2*cm
    )

    def ps(name, **kw):
        return ParagraphStyle(name, fontName="NanumGothic",
                              fontSize=10.5, leading=18,
                              textColor=colors.HexColor("#333333"), **kw)

    S = {
        "cover_title": ParagraphStyle("cv_t", fontName="NanumSquareB", fontSize=26,
            textColor=C_HEADER, alignment=TA_CENTER, spaceAfter=6),
        "cover_sub":   ParagraphStyle("cv_s", fontName="NanumGothicBold", fontSize=14,
            textColor=colors.HexColor("#555555"), alignment=TA_CENTER, spaceAfter=4),
        "h1":  ParagraphStyle("h1", fontName="NanumSquareB", fontSize=14,
            textColor=colors.white, leftIndent=8),
        "h2":  ParagraphStyle("h2", fontName="NanumGothicBold", fontSize=12,
            textColor=C_HEADER, spaceBefore=12, spaceAfter=4, leftIndent=4),
        "body":   ps("body",   leftIndent=12, spaceAfter=3),
        "bullet": ps("bullet", leftIndent=26, spaceAfter=2),
        "note":   ParagraphStyle("note", fontName="NanumGothicBold", fontSize=10,
            textColor=colors.HexColor("#7a4400"), leftIndent=16, spaceAfter=3),
        "caption":ParagraphStyle("cap", fontName="NanumGothic", fontSize=9,
            textColor=colors.HexColor("#666666"), alignment=TA_CENTER, spaceAfter=8),
        "footer": ParagraphStyle("ft",  fontName="NanumGothic", fontSize=9,
            textColor=colors.HexColor("#888888"), alignment=TA_CENTER),
        "tag": ParagraphStyle("tag", fontName="NanumGothicBold", fontSize=10,
            textColor=colors.white, alignment=TA_CENTER),
    }

    def sec_hdr(text, color):
        tbl = Table([[Paragraph(text, S["h1"])]], colWidths=[16.6*cm])
        tbl.setStyle(TableStyle([
            ("BACKGROUND",    (0,0),(-1,-1), color),
            ("TOPPADDING",    (0,0),(-1,-1), 8),
            ("BOTTOMPADDING", (0,0),(-1,-1), 8),
            ("LEFTPADDING",   (0,0),(-1,-1), 12),
        ]))
        return tbl

    def channel_tag(label, color):
        tbl = Table([[Paragraph(label, S["tag"])]], colWidths=[4*cm])
        tbl.setStyle(TableStyle([
            ("BACKGROUND",    (0,0),(-1,-1), color),
            ("TOPPADDING",    (0,0),(-1,-1), 4),
            ("BOTTOMPADDING", (0,0),(-1,-1), 4),
            ("ROUNDEDCORNERS",[4,4,4,4]),
        ]))
        return tbl

    def img_block(path, caption_text, w=16*cm):
        from reportlab.platypus import Image as RLImage
        im = RLImage(path, width=w, height=w*380/760)
        cap = Paragraph(f"▲ {caption_text}", S["caption"])
        return KeepTogether([im, cap])

    story = []

    # ── 표지
    story.append(Spacer(1, 4*cm))
    story.append(Paragraph("업무 인수인계 자료", S["cover_title"]))
    story.append(Spacer(1, 0.3*cm))
    story.append(HRFlowable(width="70%", thickness=2, color=C_HEADER, spaceAfter=10))
    story.append(Paragraph("사내 업무채널 사용 방법", S["cover_sub"]))
    story.append(Spacer(1, 1*cm))

    # 채널 태그 4개 가로 배열
    tag_row = Table(
        [[channel_tag("💬  슬랙", C_SLACK),
          channel_tag("📹  줌", C_ZOOM),
          channel_tag("🎙 클로바노트", C_CLOVA),
          channel_tag("📦  드롭박스", C_DROP)]],
        colWidths=[4.1*cm, 4.1*cm, 4.1*cm, 4.1*cm]
    )
    tag_row.setStyle(TableStyle([
        ("ALIGN",  (0,0),(-1,-1), "CENTER"),
        ("VALIGN", (0,0),(-1,-1), "MIDDLE"),
        ("LEFTPADDING",  (0,0),(-1,-1), 4),
        ("RIGHTPADDING", (0,0),(-1,-1), 4),
    ]))
    story.append(tag_row)
    story.append(Spacer(1, 0.8*cm))

    ci = Table([
        ["작성일",    "2026년 06월 16일"],
        ["작성 부서", "경영지원팀"],
        ["채널 구성", "슬랙(업무소통) / 줌(화상회의) / 클로바노트(회의 녹음·요약) / 드롭박스(파일보관)"],
    ], colWidths=[3.5*cm, 12*cm])
    ci.setStyle(TableStyle([
        ("FONTNAME",     (0,0),(-1,-1), "NanumGothic"),
        ("FONTNAME",     (0,0),(0,-1),  "NanumGothicBold"),
        ("FONTSIZE",     (0,0),(-1,-1), 11),
        ("TEXTCOLOR",    (0,0),(0,-1),  C_HEADER),
        ("VALIGN",       (0,0),(-1,-1), "TOP"),
        ("TOPPADDING",   (0,0),(-1,-1), 5),
        ("BOTTOMPADDING",(0,0),(-1,-1), 5),
        ("LINEBELOW",    (0,0),(-1,-2), 0.5, colors.HexColor("#cccccc")),
    ]))
    story.append(ci)
    story.append(PageBreak())

    # ══════════════════════════════════════════════════════════
    # 1. 슬랙
    # ══════════════════════════════════════════════════════════
    story.append(sec_hdr("1. 슬랙 (Slack)  –  업무 소통 채널", C_SLACK))
    story.append(Spacer(1, 0.3*cm))
    story.append(Paragraph(
        "슬랙은 사내 공식 업무 소통 도구입니다. "
        "PC 앱·모바일 앱·웹(slack.com) 모두 사용 가능하며, "
        "이메일보다 빠른 실시간 소통과 채널별 주제 분류가 특징입니다.",
        S["body"]))

    story.append(Paragraph("■ 채널 구조", S["h2"]))
    ch_data = [
        [Paragraph(c, ParagraphStyle("th", fontName="NanumGothicBold",
            fontSize=10, textColor=colors.white, alignment=TA_CENTER))
         for c in ["유형", "특징", "주요 채널 예시"]],
        [Paragraph("# 공개채널", S["body"]),
         Paragraph("전체 구성원 참여 가능", S["body"]),
         Paragraph("#공지사항  #전체  #업무요청", S["body"])],
        [Paragraph("🔒 비공개채널", S["body"]),
         Paragraph("초대된 멤버만 접근", S["body"]),
         Paragraph("#인사팀  #경영지원  #임원", S["body"])],
        [Paragraph("✉ 다이렉트메시지", S["body"]),
         Paragraph("1:1 또는 소그룹 비공개 대화", S["body"]),
         Paragraph("개인 간 업무 논의", S["body"])],
    ]
    ch_tbl = Table(ch_data, colWidths=[2.2*cm, 7.2*cm, 7.2*cm])
    ch_tbl.setStyle(TableStyle([
        ("BACKGROUND",    (0,0),(-1,0), C_SLACK),
        ("ROWBACKGROUNDS",(0,1),(-1,-1),[colors.white, colors.HexColor("#f5eef8")]),
        ("GRID",          (0,0),(-1,-1), 0.4, colors.HexColor("#c0a0cc")),
        ("VALIGN",        (0,0),(-1,-1), "MIDDLE"),
        ("TOPPADDING",    (0,0),(-1,-1), 5),
        ("BOTTOMPADDING", (0,0),(-1,-1), 5),
        ("LEFTPADDING",   (0,0),(-1,-1), 8),
    ]))
    story.append(ch_tbl)
    story.append(Spacer(1, 0.3*cm))

    story.append(img_block(img_slack_ch, "슬랙 채널 구조 예시"))

    story.append(Paragraph("■ 주요 기능", S["h2"]))
    for s in [
        "@ 멘션 : @이름 으로 특정인 호출 / @here 로 현재 접속자 전체 / @channel 로 채널 전체",
        "스레드 : 메시지에 [답글 달기] 클릭 → 대화 흐름을 깔끔하게 정리",
        "이모지 반응 : 메시지에 마우스 오버 → 😊 클릭으로 간단 반응",
        "파일 첨부 : 입력창 왼쪽 📎 클릭 또는 드래그 앤 드롭",
        "검색 : 상단 검색창 → 키워드·채널·기간 필터로 과거 메시지 검색",
        "북마크 : 중요 메시지에 마우스 오버 → 🔖 저장 → 나중에 확인",
    ]:
        story.append(Paragraph(f"• {s}", S["bullet"]))

    story.append(Spacer(1, 0.2*cm))
    story.append(img_block(img_slack_use, "슬랙 주요 기능 화면 예시"))
    story.append(Paragraph("※ 업무 관련 사항은 이메일보다 슬랙 채널 우선 사용", S["note"]))
    story.append(Paragraph("※ 알림 설정 : 프로필 아이콘 → 환경설정 → 알림 메뉴에서 조정 가능", S["note"]))
    story.append(PageBreak())

    # ══════════════════════════════════════════════════════════
    # 2. 줌
    # ══════════════════════════════════════════════════════════
    story.append(sec_hdr("2. 줌 (Zoom)  –  화상 회의", C_ZOOM))
    story.append(Spacer(1, 0.3*cm))
    story.append(Paragraph(
        "줌은 사내·외 화상 회의 도구입니다. "
        "PC 앱을 설치하거나 웹 브라우저로 참여할 수 있으며, "
        "회의 링크를 슬랙 채널에 공유하여 참여자를 안내합니다.",
        S["body"]))

    story.append(Paragraph("■ 회의 시작 및 예약", S["h2"]))
    for i, s in enumerate([
        "Zoom 앱 실행 후 로그인 (회사 계정 사용)",
        "즉시 시작 : [새 회의] 클릭 → 참여자에게 초대 링크 공유",
        "예약 : [예약] → 날짜·시간·반복 설정 → 초대 링크 복사 → 슬랙에 공유",
    ], 1):
        story.append(Paragraph(f"{i}. {s}", S["bullet"]))

    story.append(Paragraph("■ 회의 참여", S["h2"]))
    for i, s in enumerate([
        "슬랙에서 수신된 회의 링크 클릭 → 자동 참여",
        "또는 Zoom 앱 → [참가] → 회의 ID 입력",
    ], 1):
        story.append(Paragraph(f"{i}. {s}", S["bullet"]))

    story.append(Paragraph("■ 회의 중 주요 기능", S["h2"]))
    func_data = [
        [Paragraph(c, ParagraphStyle("th2", fontName="NanumGothicBold",
            fontSize=10, textColor=colors.white, alignment=TA_CENTER))
         for c in ["기능", "방법"]],
        [Paragraph("마이크 음소거/해제", S["body"]),
         Paragraph("하단 🎤 클릭 (단축키 : 스페이스바 누르는 동안 임시 해제)", S["body"])],
        [Paragraph("카메라 켜기/끄기", S["body"]),
         Paragraph("하단 📷 클릭", S["body"])],
        [Paragraph("화면 공유", S["body"]),
         Paragraph("하단 [공유] → 공유할 화면·창·파일 선택", S["body"])],
        [Paragraph("채팅", S["body"]),
         Paragraph("하단 [채팅] → 전체 또는 개별 참여자에게 메시지", S["body"])],
        [Paragraph("녹화", S["body"]),
         Paragraph("[녹화] 클릭 → 로컬 저장 또는 클라우드 저장 선택", S["body"])],
        [Paragraph("회의 종료", S["body"]),
         Paragraph("[종료] → 나만 나가기 / 모든 참여자 종료 선택", S["body"])],
    ]
    func_tbl = Table(func_data, colWidths=[4*cm, 12.6*cm])
    func_tbl.setStyle(TableStyle([
        ("BACKGROUND",    (0,0),(-1,0), C_ZOOM),
        ("ROWBACKGROUNDS",(0,1),(-1,-1),[colors.white, colors.HexColor("#eef2ff")]),
        ("GRID",          (0,0),(-1,-1), 0.4, colors.HexColor("#90a8ee")),
        ("VALIGN",        (0,0),(-1,-1), "MIDDLE"),
        ("TOPPADDING",    (0,0),(-1,-1), 5),
        ("BOTTOMPADDING", (0,0),(-1,-1), 5),
        ("LEFTPADDING",   (0,0),(-1,-1), 8),
    ]))
    story.append(func_tbl)
    story.append(Spacer(1, 0.3*cm))
    story.append(img_block(img_zoom, "줌 회의 화면 예시"))
    story.append(Paragraph("※ 회의 링크는 슬랙 해당 채널에 미리 공유 (회의 10분 전 권장)", S["note"]))
    story.append(Paragraph("※ 녹화 파일은 클로바노트에 업로드하여 회의록으로 변환 가능", S["note"]))
    story.append(PageBreak())

    # ══════════════════════════════════════════════════════════
    # 3. 클로바노트
    # ══════════════════════════════════════════════════════════
    story.append(sec_hdr("3. 클로바노트 (ClovaNote)  –  회의 녹음 및 요약", C_CLOVA))
    story.append(Spacer(1, 0.3*cm))
    story.append(Paragraph(
        "클로바노트는 AI 기반 회의 녹음·텍스트 변환·요약 도구입니다. "
        "모바일 앱으로 현장 녹음하거나 줌 녹화 파일을 업로드하여 "
        "자동으로 회의록을 생성합니다.",
        S["body"]))

    story.append(Paragraph("■ 현장 회의 녹음 (모바일 앱)", S["h2"]))
    for i, s in enumerate([
        "클로바노트 앱 실행 → [새 노트] 탭",
        "[녹음 시작] → 회의 진행 (자동으로 음성 인식 및 화자 분리)",
        "회의 종료 후 [정지] → 자동으로 텍스트 변환 완료",
        "변환 내용 검토 후 오류 수정",
        "AI 요약 탭에서 핵심 내용 자동 정리 확인",
        "[공유] → 링크 복사 → 슬랙 해당 채널에 공유",
    ], 1):
        story.append(Paragraph(f"{i}. {s}", S["bullet"]))

    story.append(Paragraph("■ 줌 녹화 파일 업로드 (웹)", S["h2"]))
    for i, s in enumerate([
        "clovanote.naver.com 접속 후 로그인",
        "[새 노트] → [파일 업로드] → 줌 녹화 파일(mp4/mp3) 선택",
        "업로드 완료 후 자동으로 텍스트 변환 (파일 길이에 따라 수분 소요)",
        "변환 완료 후 내용 확인 및 공유",
    ], 1):
        story.append(Paragraph(f"{i}. {s}", S["bullet"]))

    story.append(Spacer(1, 0.3*cm))
    story.append(img_block(img_clova, "클로바노트 회의 녹음·요약 화면 예시"))
    story.append(Paragraph("※ 회의록(텍스트)은 드롭박스 '회의록/YYYY-MM' 폴더에 PDF로 저장·보관", S["note"]))
    story.append(Paragraph("※ 외부 참여자가 있는 경우 사전에 녹음 동의 여부 확인 필요", S["note"]))
    story.append(PageBreak())

    # ══════════════════════════════════════════════════════════
    # 4. 드롭박스
    # ══════════════════════════════════════════════════════════
    story.append(sec_hdr("4. 드롭박스 (Dropbox)  –  파일 및 문서 보관", C_DROP))
    story.append(Spacer(1, 0.3*cm))
    story.append(Paragraph(
        "드롭박스는 사내 공식 파일 보관·공유 도구입니다. "
        "PC 앱을 설치하면 탐색기에서 바로 접근할 수 있으며, "
        "팀 폴더는 초대된 구성원 간 자동으로 동기화됩니다.",
        S["body"]))

    story.append(Paragraph("■ 폴더 구조 (예시)", S["h2"]))
    folder_data = [
        [Paragraph(c, ParagraphStyle("th3", fontName="NanumGothicBold",
            fontSize=10, textColor=colors.white, alignment=TA_CENTER))
         for c in ["폴더명", "보관 내용"]],
        [Paragraph("📁 계약서/", S["body"]),
         Paragraph("나라장터·모두싸인 계약서 PDF (연도별 하위 폴더)", S["body"])],
        [Paragraph("📁 세금계산서/", S["body"]),
         Paragraph("월별 세금계산서 PDF", S["body"])],
        [Paragraph("📁 회의록/", S["body"]),
         Paragraph("클로바노트 변환 회의록 (YYYY-MM 하위 폴더)", S["body"])],
        [Paragraph("📁 급여/", S["body"]),
         Paragraph("급여대장, 사업소득지급대장, 급여명세서", S["body"])],
        [Paragraph("📁 증명서류/", S["body"]),
         Paragraph("납세증명서, 완납증명서, 등기부등본 등 발급 서류", S["body"])],
        [Paragraph("📁 인수인계/", S["body"]),
         Paragraph("업무 인수인계 자료", S["body"])],
    ]
    folder_tbl = Table(folder_data, colWidths=[4.5*cm, 12.1*cm])
    folder_tbl.setStyle(TableStyle([
        ("BACKGROUND",    (0,0),(-1,0), C_DROP),
        ("ROWBACKGROUNDS",(0,1),(-1,-1),[colors.white, colors.HexColor("#eef0ff")]),
        ("GRID",          (0,0),(-1,-1), 0.4, colors.HexColor("#88a0ee")),
        ("VALIGN",        (0,0),(-1,-1), "MIDDLE"),
        ("TOPPADDING",    (0,0),(-1,-1), 5),
        ("BOTTOMPADDING", (0,0),(-1,-1), 5),
        ("LEFTPADDING",   (0,0),(-1,-1), 8),
    ]))
    story.append(folder_tbl)
    story.append(Spacer(1, 0.3*cm))

    story.append(Paragraph("■ 파일 업로드 및 공유", S["h2"]))
    for i, s in enumerate([
        "PC 앱 설치 시 : 탐색기의 Dropbox 폴더에 파일 복사/이동 → 자동 동기화",
        "웹(dropbox.com) : [업로드] 버튼으로 파일 직접 업로드",
        "공유 : 파일·폴더 우클릭 → [Dropbox 링크 복사] → 슬랙 등에 링크 붙여넣기",
        "팀 폴더 초대 : 폴더 우클릭 → [공유] → 이메일 입력 후 초대",
    ], 1):
        story.append(Paragraph(f"{i}. {s}", S["bullet"]))

    story.append(Spacer(1, 0.3*cm))
    story.append(img_block(img_dropbox, "드롭박스 폴더 구조 및 공유 화면 예시"))
    story.append(Paragraph("※ 중요 문서는 반드시 드롭박스에 보관 (로컬 PC에만 저장 금지)", S["note"]))
    story.append(Paragraph("※ 파일명 규칙 : YYYY-MM-DD_문서명.확장자  (예: 2026-06-16_계약서.pdf)", S["note"]))
    story.append(Paragraph("※ 팀 폴더 접근 권한 없을 경우 담당자에게 초대 요청", S["note"]))

    story.append(Spacer(1, 0.8*cm))
    story.append(HRFlowable(width="100%", thickness=1,
                             color=colors.HexColor("#cccccc"), spaceAfter=8))
    story.append(Paragraph(
        "본 자료는 인수인계 목적으로 작성되었습니다. 채널 정책 변경 시 담당자가 업데이트하십시오.",
        S["footer"]))

    doc.build(story)
    print(f"PDF 생성 완료 : {out_path}")


# ══════════════════════════════════════════════════════════════
# DOCX 생성
# ══════════════════════════════════════════════════════════════

def set_cell_bg(cell, hex_color):
    tc = cell._tc; tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color.lstrip("#"))
    tcPr.append(shd)

def ko(run):
    run.font.name = "나눔고딕"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "나눔고딕")

def add_sec_hdr(doc, text, bg_hex):
    tbl = doc.add_table(rows=1, cols=1); tbl.style = "Table Grid"
    cell = tbl.rows[0].cells[0]; cell.text = text
    set_cell_bg(cell, bg_hex)
    p = cell.paragraphs[0]; r = p.runs[0]; ko(r)
    r.font.bold = True; r.font.size = Pt(14)
    r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    doc.add_paragraph()

def add_para(doc, text, bold=False, size=10.5, indent=0.3, color=None):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(indent)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(text); ko(r)
    r.font.size = Pt(size); r.font.bold = bold
    if color: r.font.color.rgb = color
    return p

def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Cm(0.8)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(text); ko(r); r.font.size = Pt(10.5)

def add_img(doc, path, caption, width_cm=14.5):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(path, width=Cm(width_cm))
    cap = doc.add_paragraph(f"▲ {caption}")
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in cap.runs:
        ko(r); r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(0x88, 0x88, 0x88)
    doc.add_paragraph()

def note(doc, text):
    add_para(doc, text, color=RGBColor(0x7a, 0x44, 0x00))

def simple_table(doc, headers, rows, hdr_hex, col_widths):
    tbl = doc.add_table(rows=1+len(rows), cols=len(headers))
    tbl.style = "Table Grid"
    for ci, h in enumerate(headers):
        c = tbl.rows[0].cells[ci]; c.text = h
        set_cell_bg(c, hdr_hex)
        r = c.paragraphs[0].runs[0]; ko(r)
        r.font.bold = True; r.font.size = Pt(10)
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            c = tbl.rows[ri+1].cells[ci]; c.text = val
            set_cell_bg(c, "FFFFFF" if ri%2==0 else "EEF0FF")
            r = c.paragraphs[0].runs[0]; ko(r); r.font.size = Pt(10)
    for ri in range(1+len(rows)):
        for ci, w in enumerate(col_widths):
            tbl.rows[ri].cells[ci].width = Cm(w)
    doc.add_paragraph()


def build_docx(out_path):
    doc = Document()
    for sec in doc.sections:
        sec.page_width=Cm(21); sec.page_height=Cm(29.7)
        sec.top_margin=Cm(2); sec.bottom_margin=Cm(2)
        sec.left_margin=Cm(2.5); sec.right_margin=Cm(2.5)

    # 표지
    for _ in range(5):
        doc.add_paragraph()
    tp = doc.add_paragraph(); tp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tr = tp.add_run("업무 인수인계 자료"); ko(tr)
    tr.font.size=Pt(26); tr.font.bold=True
    tr.font.color.rgb=RGBColor(0x2d,0x2d,0x2d)

    sp = doc.add_paragraph(); sp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sr = sp.add_run("사내 업무채널 사용 방법"); ko(sr)
    sr.font.size=Pt(15); sr.font.bold=True
    sr.font.color.rgb=RGBColor(0x55,0x55,0x55)
    doc.add_paragraph()

    ci_tbl = doc.add_table(rows=3, cols=2); ci_tbl.style="Table Grid"
    for ri,(k,v) in enumerate([
        ("작성일","2026년 06월 16일"),
        ("작성 부서","경영지원팀"),
        ("채널 구성","슬랙(업무소통) / 줌(화상회의) / 클로바노트(회의 녹음·요약) / 드롭박스(파일보관)"),
    ]):
        cells=ci_tbl.rows[ri].cells
        set_cell_bg(cells[0],"dce6f8"); cells[0].text=k; cells[1].text=v
        for ci_,c in enumerate(cells):
            r=c.paragraphs[0].runs[0]; ko(r); r.font.size=Pt(10.5)
            if ci_==0: r.font.bold=True
        cells[0].width=Cm(3.5); cells[1].width=Cm(12)
    doc.add_page_break()

    # 1. 슬랙
    add_sec_hdr(doc, "1. 슬랙 (Slack)  –  업무 소통 채널", "4a154b")
    add_para(doc,
        "슬랙은 사내 공식 업무 소통 도구입니다. PC 앱·모바일 앱·웹 모두 사용 가능하며, "
        "채널별로 주제를 분류하여 이메일보다 빠르게 소통합니다.")
    add_para(doc, "■ 채널 구조", bold=True)
    simple_table(doc,
        ["유형","특징","주요 채널 예시"],
        [["# 공개채널","전체 구성원 참여 가능","#공지사항  #전체  #업무요청"],
         ["🔒 비공개채널","초대된 멤버만 접근","#인사팀  #경영지원  #임원"],
         ["✉ 다이렉트메시지","1:1 또는 소그룹 비공개","개인 간 업무 논의"]],
        "4a154b", [3.5,6,7])
    add_para(doc, "■ 주요 기능", bold=True)
    for s in [
        "@ 멘션 : @이름 으로 특정인 호출 / @here 현재 접속자 전체 / @channel 채널 전체",
        "스레드 : 메시지에 [답글 달기] 클릭 → 대화 흐름 정리",
        "파일 첨부 : 입력창 왼쪽 📎 클릭 또는 드래그 앤 드롭",
        "검색 : 상단 검색창 → 키워드·채널·기간 필터",
        "북마크 : 중요 메시지 오버 → 🔖 저장",
    ]:
        add_bullet(doc, f"• {s}")
    doc.add_paragraph()
    add_img(doc, img_slack_ch,  "슬랙 채널 구조 예시")
    add_img(doc, img_slack_use, "슬랙 주요 기능 화면 예시")
    note(doc, "※ 업무 관련 사항은 이메일보다 슬랙 채널 우선 사용")
    note(doc, "※ 알림 설정 : 프로필 → 환경설정 → 알림 메뉴에서 조정")
    doc.add_page_break()

    # 2. 줌
    add_sec_hdr(doc, "2. 줌 (Zoom)  –  화상 회의", "0b5cff")
    add_para(doc,
        "줌은 사내·외 화상 회의 도구입니다. PC 앱 또는 웹 브라우저로 참여하며, "
        "회의 링크를 슬랙 채널에 공유하여 참여자를 안내합니다.")
    add_para(doc, "■ 회의 시작 및 예약", bold=True)
    for i,s in enumerate([
        "Zoom 앱 실행 후 회사 계정으로 로그인",
        "즉시 시작 : [새 회의] 클릭 → 참여자에게 초대 링크 공유",
        "예약 : [예약] → 날짜·시간 설정 → 초대 링크 복사 → 슬랙에 공유",
    ],1):
        add_bullet(doc, f"{i}. {s}")
    add_para(doc, "■ 회의 중 주요 기능", bold=True)
    simple_table(doc,
        ["기능","방법"],
        [["마이크 음소거/해제","하단 🎤 클릭 (스페이스바 누르는 동안 임시 해제)"],
         ["카메라 켜기/끄기","하단 📷 클릭"],
         ["화면 공유","하단 [공유] → 공유할 화면·창 선택"],
         ["채팅","하단 [채팅] → 전체 또는 개별 참여자에게 메시지"],
         ["녹화","[녹화] 클릭 → 로컬 또는 클라우드 저장"],
         ["회의 종료","[종료] → 나만 나가기 / 모든 참여자 종료"]],
        "0b5cff", [4,13])
    add_img(doc, img_zoom, "줌 회의 화면 예시")
    note(doc, "※ 회의 링크는 슬랙 해당 채널에 회의 10분 전 미리 공유 권장")
    note(doc, "※ 녹화 파일은 클로바노트에 업로드하여 회의록으로 변환 가능")
    doc.add_page_break()

    # 3. 클로바노트
    add_sec_hdr(doc, "3. 클로바노트 (ClovaNote)  –  회의 녹음 및 요약", "03c75a")
    add_para(doc,
        "클로바노트는 AI 기반 회의 녹음·텍스트 변환·요약 도구입니다. "
        "모바일 앱으로 현장 녹음하거나 줌 녹화 파일을 업로드하여 자동으로 회의록을 생성합니다.")
    add_para(doc, "■ 현장 회의 녹음 (모바일 앱)", bold=True)
    for i,s in enumerate([
        "클로바노트 앱 실행 → [새 노트] 탭",
        "[녹음 시작] → 회의 진행 (자동 음성 인식 및 화자 분리)",
        "회의 종료 후 [정지] → 자동 텍스트 변환 완료",
        "변환 내용 검토 후 오류 수정",
        "AI 요약 탭에서 핵심 내용 확인",
        "[공유] → 링크 복사 → 슬랙 해당 채널에 공유",
    ],1):
        add_bullet(doc, f"{i}. {s}")
    add_para(doc, "■ 줌 녹화 파일 업로드 (웹)", bold=True)
    for i,s in enumerate([
        "clovanote.naver.com 접속 후 로그인",
        "[새 노트] → [파일 업로드] → 줌 녹화 파일(mp4/mp3) 선택",
        "업로드 완료 후 자동 텍스트 변환 (파일 길이에 따라 수분 소요)",
        "변환 완료 후 내용 확인 및 공유",
    ],1):
        add_bullet(doc, f"{i}. {s}")
    doc.add_paragraph()
    add_img(doc, img_clova, "클로바노트 회의 녹음·요약 화면 예시")
    note(doc, "※ 회의록은 드롭박스 '회의록/YYYY-MM' 폴더에 PDF로 저장·보관")
    note(doc, "※ 외부 참여자가 있는 경우 사전에 녹음 동의 여부 확인 필요")
    doc.add_page_break()

    # 4. 드롭박스
    add_sec_hdr(doc, "4. 드롭박스 (Dropbox)  –  파일 및 문서 보관", "0061ff")
    add_para(doc,
        "드롭박스는 사내 공식 파일 보관·공유 도구입니다. "
        "PC 앱을 설치하면 탐색기에서 바로 접근할 수 있으며, "
        "팀 폴더는 초대된 구성원 간 자동으로 동기화됩니다.")
    add_para(doc, "■ 폴더 구조 (예시)", bold=True)
    simple_table(doc,
        ["폴더명","보관 내용"],
        [["📁 계약서/","나라장터·모두싸인 계약서 PDF (연도별 하위 폴더)"],
         ["📁 세금계산서/","월별 세금계산서 PDF"],
         ["📁 회의록/","클로바노트 변환 회의록 (YYYY-MM 하위 폴더)"],
         ["📁 급여/","급여대장, 사업소득지급대장, 급여명세서"],
         ["📁 증명서류/","납세증명서, 완납증명서, 등기부등본 등"],
         ["📁 인수인계/","업무 인수인계 자료"]],
        "0061ff", [4.5,12])
    add_para(doc, "■ 파일 업로드 및 공유", bold=True)
    for i,s in enumerate([
        "PC 앱 : 탐색기의 Dropbox 폴더에 파일 복사/이동 → 자동 동기화",
        "웹(dropbox.com) : [업로드] 버튼으로 직접 업로드",
        "공유 : 파일·폴더 우클릭 → [Dropbox 링크 복사] → 슬랙에 붙여넣기",
        "팀 폴더 초대 : 폴더 우클릭 → [공유] → 이메일 입력 후 초대",
    ],1):
        add_bullet(doc, f"{i}. {s}")
    doc.add_paragraph()
    add_img(doc, img_dropbox, "드롭박스 폴더 구조 및 공유 화면 예시")
    note(doc, "※ 중요 문서는 반드시 드롭박스에 보관 (로컬 PC에만 저장 금지)")
    note(doc, "※ 파일명 규칙 : YYYY-MM-DD_문서명.확장자  (예: 2026-06-16_계약서.pdf)")
    note(doc, "※ 팀 폴더 접근 권한 없을 경우 담당자에게 초대 요청")

    doc.add_paragraph()
    fp = doc.add_paragraph(
        "본 자료는 인수인계 목적으로 작성되었습니다. 채널 정책 변경 시 담당자가 업데이트하십시오.")
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in fp.runs:
        ko(r); r.font.size=Pt(9)
        r.font.color.rgb=RGBColor(0x88,0x88,0x88)

    doc.save(out_path)
    print(f"DOCX 생성 완료 : {out_path}")


if __name__ == "__main__":
    pdf_path  = os.path.join(OUT_DIR, "사내업무채널_업무인수인계.pdf")
    docx_path = os.path.join(OUT_DIR, "사내업무채널_업무인수인계.docx")
    build_pdf(pdf_path)
    build_docx(docx_path)
    print("모든 파일 생성 완료.")
